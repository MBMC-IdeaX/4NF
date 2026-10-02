"""
Appends the closing slide to an existing Bhada-slides.pptx.

Separate from make-deck.py on purpose. Once a human has typed the team names
into the deck, regenerating it throws that work away — so this opens the file
that is already there, adds one slide at the end, and saves it back. The same
slide is also defined in make-deck.py, so a full regeneration produces it too.

    python scripts/add-closing-slide.py

Safe to run twice: it checks for the slide it adds and does nothing if it is
already present.
"""

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
DECK = ROOT / "Bhada-slides.pptx"
QR = ROOT / "qr-bhada.png"

INK = RGBColor(0x0F, 0x0E, 0x0C)
PANEL = RGBColor(0x17, 0x15, 0x0F)
RULE = RGBColor(0x2B, 0x27, 0x21)
RULE_HI = RGBColor(0x45, 0x3F, 0x34)
READ = RGBColor(0xEC, 0xE7, 0xD8)
DIM = RGBColor(0x8E, 0x87, 0x79)
PLATE = RGBColor(0xA8, 0x20, 0x2F)
LIVE = RGBColor(0x6F, 0xBF, 0x73)

DISPLAY = "Bahnschrift SemiBold Condensed"
BODY = "Segoe UI"
DEVA = "Nirmala UI"
MONO = "Consolas"

MARKER = "Any questions?"


def rect(s, x, y, w, h, fill=None, line=None):
    shape = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.75)
    shape.shadow.inherit = False
    return shape


def text(s, x, y, w, h, runs, align=PP_ALIGN.LEFT, spacing=None):
    box = s.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraphs = runs if isinstance(runs[0], list) else [runs]
    for i, para in enumerate(paragraphs):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = spacing
        for content, size, colour, font, bold in para:
            r = p.add_run()
            r.text = content
            r.font.size = Pt(size)
            r.font.color.rgb = colour
            r.font.name = font
            r.font.bold = bold
    return box


def build_closing_slide(prs):
    """The last thing on the wall while they are asking questions.

    It stays up for as long as the Q&A runs, which is usually longer than the
    pitch itself — so it carries the three things worth being asked about, and a
    code a judge can scan from their seat to open the live thing. A slide that
    only says THANK YOU wastes the most valuable minute in the room.
    """
    W, H = prs.slide_width, prs.slide_height
    M = Inches(0.62)

    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = rect(s, 0, 0, W, H, fill=INK)
    del bg

    rect(s, M, M + Inches(0.06), Inches(0.05), Inches(0.16), fill=PLATE)
    text(s, M + Inches(0.18), M, Inches(9), Inches(0.3),
         [("THANK YOU", 11, DIM, BODY, False)])

    text(s, M, Inches(1.15), Inches(8.4), Inches(2.2),
         [[("प्रश्न ?", 72, READ, DEVA, True)],
          [(MARKER, 46, PLATE, DISPLAY, True)]])

    # Three invitations. Naming what you want to be asked is how you steer a
    # Q&A towards the parts of the work that are strongest.
    prompts = [
        ("Try it now", "open it and turn your wifi off", LIVE),
        ("Ask about the door", "the fare machine is the load limiter", "amber"),
        ("Ask what a route earns", "nobody in Nepal has that number", READ),
    ]
    y = Inches(3.75)
    for i, (title, sub, colour) in enumerate(prompts):
        yy = y + i * Inches(0.86)
        bar = LIVE if colour is LIVE else (RGBColor(0xE0, 0xA4, 0x37) if colour == "amber" else READ)
        rect(s, M, yy, Inches(0.06), Inches(0.58), fill=bar)
        text(s, M + Inches(0.32), yy - Inches(0.02), Inches(6.2), Inches(0.35),
             [(title, 22, READ, DISPLAY, True)])
        text(s, M + Inches(0.32), yy + Inches(0.33), Inches(6.2), Inches(0.3),
             [(sub, 12, DIM, BODY, False)])

    # The live link, as a code they can scan from their seat.
    panel_x = Inches(8.3)
    rect(s, panel_x, Inches(1.15), Inches(4.42), Inches(4.6), fill=PANEL, line=RULE_HI)
    if QR.exists():
        s.shapes.add_picture(str(QR), panel_x + Inches(1.06), Inches(1.5),
                             width=Inches(2.3), height=Inches(2.3))
    text(s, panel_x, Inches(4.0), Inches(4.42), Inches(0.4),
         [[("bhada-one.vercel.app", 19, LIVE, MONO, True)]], align=PP_ALIGN.CENTER)
    text(s, panel_x + Inches(0.4), Inches(4.5), Inches(3.62), Inches(1.0),
         [[("Scan it. Then put your phone in airplane mode and use it anyway.",
            13, DIM, BODY, False)]], align=PP_ALIGN.CENTER)

    rect(s, M, H - Inches(0.72), W - 2 * M, Emu(9525), fill=RULE)
    text(s, M, H - Inches(0.6), Inches(9), Inches(0.3),
         [("Bhada · भाडा · Ratna Park — Koteshwor, Kathmandu", 9.5, DIM, BODY, False)])
    return s


def already_present(prs):
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame and MARKER in shape.text_frame.text:
                return True
    return False


def main():
    if not DECK.exists():
        print(f"{DECK.name} not found. Run make-deck.py first.")
        return 1

    prs = Presentation(str(DECK))
    if already_present(prs):
        print("closing slide already there — nothing to do")
        return 0

    before = len(prs.slides._sldIdLst)
    build_closing_slide(prs)
    try:
        prs.save(str(DECK))
    except PermissionError:
        print(f"{DECK.name} is open in PowerPoint. Close it and run this again.")
        return 1
    print(f"added the closing slide — {before} slides in, {before + 1} out")
    return 0


if __name__ == "__main__":
    sys.exit(main())
